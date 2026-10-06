"""
FedCrack-Net XAI: Microsoft Word (.docx) Inspection Report Exporter
Generates formal editable inspection documents using python-docx,
with formatted typography, inspection tables, and embedded visual plates.
"""

import io
from typing import Dict, Any, Optional
import numpy as np
from PIL import Image as PILImage

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn


class DOCXReportExporter:
    """Compiles structured inspection reports into editable .docx documents."""

    def __init__(self):
        self.primary_color = RGBColor(30, 58, 138)    # #1E3A8A
        self.secondary_color = RGBColor(37, 99, 235)  # #2563EB
        self.text_dark = RGBColor(15, 23, 42)         # #0F172A
        self.text_muted = RGBColor(100, 116, 139)     # #64748B

    def _set_cell_background(self, cell, hex_color: str):
        """Set shading background on a table cell."""
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def _convert_numpy_to_temp_img(self, img_arr: np.ndarray) -> io.BytesIO:
        """Convert numpy image array to in-memory JPEG stream for docx insertion."""
        pil_img = PILImage.fromarray(img_arr)
        buf = io.BytesIO()
        pil_img.save(buf, format="JPEG", quality=85)
        buf.seek(0)
        return buf

    def build_docx(
        self,
        report_data: Dict[str, Any],
        original_img: Optional[np.ndarray] = None,
        detection_img: Optional[np.ndarray] = None,
        gradcam_img: Optional[np.ndarray] = None
    ) -> bytes:
        """Generate DOCX document binary buffer from report data and visual images."""
        doc = Document()

        # Set standard margins (0.75 in)
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(0.75)
            section.bottom_margin = Inches(0.75)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)

            # Footer
            footer = section.footer
            f_p = footer.paragraphs[0]
            f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            f_run = f_p.add_run("FedCrack-Net XAI | Confidential Civil Engineering Assessment")
            f_run.font.size = Pt(8)
            f_run.font.color.rgb = self.text_muted

        # 1. Title Banner
        title_p = doc.add_paragraph()
        t_run = title_p.add_run("AI Structural Inspection Report")
        t_run.font.name = "Arial"
        t_run.font.size = Pt(20)
        t_run.font.bold = True
        t_run.font.color.rgb = self.primary_color
        title_p.paragraph_format.space_after = Pt(2)

        sub_p = doc.add_paragraph()
        s_run = sub_p.add_run("FedCrack-Net XAI — Federated Explainable Structural Crack Detection System")
        s_run.font.name = "Arial"
        s_run.font.size = Pt(10)
        s_run.font.color.rgb = self.text_muted
        sub_p.paragraph_format.space_after = Pt(12)

        # 2. Inspection Information Table
        info = report_data.get("inspection_info", {})
        info_table = doc.add_table(rows=3, cols=3)
        info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        info_data = [
            [f"Date: {info.get('date', 'N/A')}", f"Time: {info.get('time', 'N/A')}", f"Inspector: {info.get('inspector', 'AI Assistant')}"],
            [f"Building: {info.get('building_name', 'N/A')}", f"Location: {info.get('location', 'N/A')}", f"Client: {info.get('client', 'N/A')}"],
            [f"Image File: {info.get('image_name', 'N/A')}", f"Project Version: {report_data.get('project_version', 'v1.0')}", "Assessment: Visual XAI"],
        ]
        for r_idx, row in enumerate(info_table.rows):
            for c_idx, cell in enumerate(row.cells):
                self._set_cell_background(cell, "F8FAFC")
                cell.text = info_data[r_idx][c_idx]
                p = cell.paragraphs[0]
                p.runs[0].font.size = Pt(8.5)
                p.runs[0].font.color.rgb = self.text_dark
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

        # Visual Plates Table
        visual_imgs = []
        if original_img is not None: visual_imgs.append(("1. Input Photo", self._convert_numpy_to_temp_img(original_img)))
        if detection_img is not None: visual_imgs.append(("2. YOLOv8 Detections", self._convert_numpy_to_temp_img(detection_img)))
        if gradcam_img is not None: visual_imgs.append(("3. Grad-CAM Saliency", self._convert_numpy_to_temp_img(gradcam_img)))

        if visual_imgs:
            h_vis = doc.add_heading(level=2)
            h_run = h_vis.add_run("Inspection Visual Plates")
            h_run.font.color.rgb = self.primary_color
            h_vis.paragraph_format.space_after = Pt(4)

            plate_table = doc.add_table(rows=2, cols=len(visual_imgs))
            plate_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for col_idx, (caption, img_buf) in enumerate(visual_imgs):
                cell_img = plate_table.cell(0, col_idx)
                p_img = cell_img.paragraphs[0]
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.add_run().add_picture(img_buf, width=Inches(2.1))

                cell_cap = plate_table.cell(1, col_idx)
                p_cap = cell_cap.paragraphs[0]
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                c_run = p_cap.add_run(caption)
                c_run.font.size = Pt(8)
                c_run.font.italic = True
                c_run.font.color.rgb = self.text_muted
            doc.add_paragraph().paragraph_format.space_after = Pt(6)

        # 3. Detection Summary Table
        summary = report_data.get("detection_summary", {})
        h_sum = doc.add_heading(level=2)
        h_run = h_sum.add_run("3. Empirical Detection & Geometric Summary")
        h_run.font.color.rgb = self.primary_color
        h_sum.paragraph_format.space_after = Pt(4)

        sum_rows = [
            ("Cracks Identified", str(summary.get("num_cracks", 0)), "Primary Crack Type", str(summary.get("crack_type", "N/A"))),
            ("Severity Tier", str(summary.get("severity", "N/A")), "Detection Confidence", f"{summary.get('confidence', 0)*100:.1f}%"),
            ("Fissure Length", f"{summary.get('length_px', 0):.1f} px", "Average Aperture", f"{summary.get('avg_width_px', 0):.1f} px"),
            ("Maximum Aperture", f"{summary.get('max_width_px', 0):.1f} px", "Orientation Angle", f"{summary.get('orientation_deg', 0):.1f}°"),
            ("Curvature Index", f"{summary.get('curvature', 1.0):.2f}", "Surface Coverage", f"{summary.get('area_coverage_pct', 0):.2f}%"),
        ]

        metric_table = doc.add_table(rows=1 + len(sum_rows), cols=4)
        metric_table.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Table Header
        headers = ["Parameter", "Measured Value", "Parameter", "Measured Value"]
        for c_idx, text in enumerate(headers):
            cell = metric_table.cell(0, c_idx)
            self._set_cell_background(cell, "1E3A8A")
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(8.5)

        for r_idx, row_vals in enumerate(sum_rows):
            bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
            for c_idx, text in enumerate(row_vals):
                cell = metric_table.cell(r_idx + 1, c_idx)
                self._set_cell_background(cell, bg)
                p = cell.paragraphs[0]
                run = p.add_run(text)
                run.font.size = Pt(8.5)
                run.font.color.rgb = self.text_dark
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

        # 4. AI Structural Assessment
        h4 = doc.add_heading(level=2)
        h4.add_run("4. AI Structural Assessment").font.color.rgb = self.primary_color
        p4 = doc.add_paragraph(report_data.get("structural_assessment", ""))
        p4.paragraph_format.line_spacing = 1.15

        # 5. Possible Engineering Causes (Visual Assessment)
        h5 = doc.add_heading(level=2)
        h5.add_run("5. Possible Engineering Causes (Visual Assessment)").font.color.rgb = self.primary_color
        for line in report_data.get("possible_causes", "").split("\n"):
            line = line.strip()
            if line:
                p = doc.add_paragraph(line)
                p.paragraph_format.line_spacing = 1.15

        # 6. Risk Assessment
        risk_info = report_data.get("risk_assessment", {})
        h6 = doc.add_heading(level=2)
        h6.add_run(f"6. Risk Assessment ({risk_info.get('risk_level', 'Moderate').upper()})").font.color.rgb = self.primary_color
        for line in risk_info.get("details", "").split("\n"):
            line = line.strip()
            if line:
                p = doc.add_paragraph(line)
                p.paragraph_format.line_spacing = 1.15

        # 7. Recommended Actions
        actions = report_data.get("recommended_actions", {})
        h7 = doc.add_heading(level=2)
        h7.add_run("7. Recommended Engineering Actions").font.color.rgb = self.primary_color

        act_table = doc.add_table(rows=4, cols=2)
        act_headers = ["Phase", "Remedial Specification"]
        for c_idx, text in enumerate(act_headers):
            cell = act_table.cell(0, c_idx)
            self._set_cell_background(cell, "1E3A8A")
            run = cell.paragraphs[0].add_run(text)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(8.5)

        act_rows = [
            ("Immediate", actions.get("immediate", "Monitor baseline")),
            ("Short-Term", actions.get("short_term", "Install telltales and seal")),
            ("Long-Term", actions.get("long_term", "Perform structural remediation")),
        ]
        for r_idx, (phase, rec) in enumerate(act_rows):
            bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
            for c_idx, text in enumerate([phase, rec]):
                cell = act_table.cell(r_idx + 1, c_idx)
                self._set_cell_background(cell, bg)
                run = cell.paragraphs[0].add_run(text)
                if c_idx == 0: run.font.bold = True
                run.font.size = Pt(8.5)
                run.font.color.rgb = self.text_dark
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

        # 8. Explainable AI (Grad-CAM)
        h8 = doc.add_heading(level=2)
        h8.add_run("8. Explainable AI (Grad-CAM) Verification").font.color.rgb = self.primary_color
        p8 = doc.add_paragraph(report_data.get("xai_explanation", ""))
        p8.paragraph_format.line_spacing = 1.15

        # 9. Engineering Notes
        h9 = doc.add_heading(level=2)
        h9.add_run("9. Engineering Notes").font.color.rgb = self.primary_color
        for line in report_data.get("engineering_notes", "").split("\n"):
            line = line.strip()
            if line:
                p = doc.add_paragraph(line)
                p.paragraph_format.line_spacing = 1.15

        # 10. Regulatory Disclaimer
        h10 = doc.add_heading(level=2)
        h10.add_run("10. Regulatory Disclaimer").font.color.rgb = self.primary_color
        p10 = doc.add_paragraph(report_data.get("disclaimer", ""))
        p10.runs[0].font.italic = True
        p10.runs[0].font.size = Pt(8.5)
        p10.runs[0].font.color.rgb = self.text_muted

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()
